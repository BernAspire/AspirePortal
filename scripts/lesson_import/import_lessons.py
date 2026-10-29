#!/usr/bin/env python3
"""Imports lesson files (PowerPoint slides, Word/PDF worksheets) into the portal.

For each lesson (files grouped by level, subject and week from their names):
  1. convert to PDF with LibreOffice,
  2. render every page to a WebP image,
  3. for slides: hide live-class-only slides and mark worked solutions,
  4. upload the PDF and page images to the private `lesson-files` bucket,
  5. create or update the course, lesson and lesson parts.

Usage:
  # Build only, into a folder you can inspect (no network):
  python3 import_lessons.py build --out ./out FILE [FILE ...]

  # Import everything in the `lesson-inbox` bucket for one centre (the GitHub job does this):
  SUPABASE_URL=... SUPABASE_SERVICE_ROLE_KEY=... \
    python3 import_lessons.py inbox --centre aspire-hub [--draft]

The service-role key is used only here, in a job, never in the app (CLAUDE.md rule 5).
Every row written carries the centre's id explicitly.
"""

from __future__ import annotations

import argparse
import json
import os
import shutil
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.parse
import urllib.request
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

from naming import PART_SLIDES, PART_WORKSHEET, LessonFile, page_title, parse_name, slide_role

SLIDE_WIDTH_PX = 1600
WORKSHEET_WIDTH_PX = 1240
WEBP_QUALITY = 82

PART_SETTINGS = {
    PART_SLIDES: {"position": 1, "title": "Lesson slides", "downloadable": False},
    PART_WORKSHEET: {"position": 2, "title": "Worksheet", "downloadable": True},
}


# --------------------------------------------------------------------------------------
# Building: files on disk → PDF + page images + page list
# --------------------------------------------------------------------------------------
@dataclass
class BuiltPage:
    image: Path
    role: str  # "page" | "solution"
    title: str


@dataclass
class BuiltPart:
    part: str
    source: Path
    pdf: Path
    pages: list[BuiltPage] = field(default_factory=list)
    hidden: list[str] = field(default_factory=list)


def to_pdf(source: Path, workdir: Path) -> Path:
    if source.suffix.lower() == ".pdf":
        return source
    subprocess.run(
        ["soffice", "--headless", "--norestore", "--convert-to", "pdf", "--outdir", str(workdir), str(source)],
        check=True,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        timeout=600,
    )
    pdf = workdir / (source.stem + ".pdf")
    if not pdf.exists():
        raise RuntimeError(f"LibreOffice did not produce a PDF for {source.name}")
    return pdf


def pdf_page_count(pdf: Path) -> int:
    out = subprocess.run(["pdfinfo", str(pdf)], check=True, capture_output=True, text=True).stdout
    for line in out.splitlines():
        if line.startswith("Pages:"):
            return int(line.split()[1])
    raise RuntimeError(f"Could not read page count of {pdf.name}")


def render_pages(pdf: Path, width_px: int, outdir: Path) -> list[Path]:
    from PIL import Image  # imported here so `naming` tests need no Pillow

    outdir.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        ["pdftoppm", "-png", "-scale-to-x", str(width_px), "-scale-to-y", "-1", str(pdf), str(outdir / "p")],
        check=True,
        stderr=subprocess.DEVNULL,
        timeout=900,
    )
    images = []
    for png in sorted(outdir.glob("p-*.png"), key=lambda p: int(p.stem.split("-")[-1])):
        webp = png.with_suffix(".webp")
        with Image.open(png) as im:
            im.convert("RGB").save(webp, "WEBP", quality=WEBP_QUALITY, method=6)
        png.unlink()
        images.append(webp)
    return images


def slide_texts(pptx: Path) -> list[tuple[bool, str, str]]:
    """(visible, title, all text) for each slide, in order."""
    from pptx import Presentation

    result = []
    for slide in Presentation(str(pptx)).slides:
        visible = slide._element.get("show") != "0"
        texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text
        if not title.strip() and texts:
            title = texts[0]
        result.append((visible, page_title(title), "\n".join(texts)))
    return result


def build_part(file: LessonFile, source: Path, workdir: Path) -> BuiltPart:
    workdir.mkdir(parents=True, exist_ok=True)
    pdf = to_pdf(source, workdir)
    width = SLIDE_WIDTH_PX if file.part == PART_SLIDES else WORKSHEET_WIDTH_PX
    images = render_pages(pdf, width, workdir / "pages")
    built = BuiltPart(part=file.part, source=source, pdf=pdf)

    if file.part == PART_SLIDES and source.suffix.lower() == ".pptx":
        # LibreOffice leaves out slides already hidden in PowerPoint; match the rest 1:1.
        slides = [s for s in slide_texts(source) if s[0]]
        if len(slides) != len(images):
            raise RuntimeError(f"{source.name}: {len(slides)} visible slides but {len(images)} PDF pages")
        for image, (_, title, _) in zip(images, slides):
            role = slide_role(title)
            if role == "hidden":
                built.hidden.append(title)
                image.unlink()
                continue
            built.pages.append(BuiltPage(image=image, role=role, title=title))
    else:
        for n, image in enumerate(images, start=1):
            built.pages.append(BuiltPage(image=image, role="page", title=f"Page {n}"))

    if not built.pages:
        raise RuntimeError(f"{source.name}: no pages left to show")
    return built


def group_files(paths: list[Path]) -> tuple[dict[tuple, dict[str, tuple[LessonFile, Path]]], list[str]]:
    """Groups files into lessons by name; returns (lessons, problems)."""
    lessons: dict[tuple, dict[str, tuple[LessonFile, Path]]] = defaultdict(dict)
    problems = []
    for path in paths:
        info = parse_name(path.name)
        if info is None:
            problems.append(f"Skipped {path.name}: name doesn't follow 'P5 - Mathematics - Worksheet 01 - Title.pptx'")
            continue
        if info.part in lessons[info.key]:
            problems.append(f"Skipped {path.name}: another {info.part} file for the same lesson")
            continue
        lessons[info.key][info.part] = (info, path)
    return lessons, problems


# --------------------------------------------------------------------------------------
# Publishing: Supabase Storage + REST with the service-role key
# --------------------------------------------------------------------------------------
class Supabase:
    def __init__(self, url: str, key: str):
        self.url = url.rstrip("/")
        self.key = key

    def _request(self, method: str, path: str, body: bytes | None = None, headers: dict | None = None) -> bytes:
        req = urllib.request.Request(self.url + path, data=body, method=method)
        req.add_header("apikey", self.key)
        req.add_header("Authorization", f"Bearer {self.key}")
        for k, v in (headers or {}).items():
            req.add_header(k, v)
        for attempt in range(4):
            try:
                with urllib.request.urlopen(req, timeout=120) as res:
                    return res.read()
            except urllib.error.HTTPError as err:
                detail = err.read().decode("utf-8", "replace")[:500]
                if err.code >= 500 and attempt < 3:
                    time.sleep(2**attempt)
                    continue
                raise RuntimeError(f"{method} {path.split('?')[0]} failed: HTTP {err.code} {detail}") from None
            except urllib.error.URLError:
                if attempt < 3:
                    time.sleep(2**attempt)
                    continue
                raise
        raise RuntimeError("unreachable")

    # REST
    def select(self, table: str, query: dict) -> list[dict]:
        qs = urllib.parse.urlencode(query)
        return json.loads(self._request("GET", f"/rest/v1/{table}?{qs}"))

    def insert(self, table: str, row: dict) -> dict:
        out = self._request(
            "POST", f"/rest/v1/{table}", json.dumps(row).encode(),
            {"Content-Type": "application/json", "Prefer": "return=representation"},
        )
        return json.loads(out)[0]

    def update(self, table: str, match: dict, values: dict) -> list[dict]:
        qs = urllib.parse.urlencode({k: f"eq.{v}" for k, v in match.items()})
        out = self._request(
            "PATCH", f"/rest/v1/{table}?{qs}", json.dumps(values).encode(),
            {"Content-Type": "application/json", "Prefer": "return=representation"},
        )
        return json.loads(out)

    # Storage
    def upload(self, bucket: str, path: str, file: Path, content_type: str) -> None:
        self._request(
            "POST", f"/storage/v1/object/{bucket}/{urllib.parse.quote(path)}", file.read_bytes(),
            {"Content-Type": content_type, "x-upsert": "true", "Cache-Control": "max-age=3600"},
        )

    def list_inbox(self, bucket: str) -> list[str]:
        body = json.dumps({"prefix": "", "limit": 1000, "sortBy": {"column": "name", "order": "asc"}}).encode()
        items = json.loads(self._request("POST", f"/storage/v1/object/list/{bucket}", body, {"Content-Type": "application/json"}))
        return [i["name"] for i in items if i.get("id")]  # folders have no id

    def download(self, bucket: str, path: str, dest: Path) -> Path:
        dest.write_bytes(self._request("GET", f"/storage/v1/object/{bucket}/{urllib.parse.quote(path)}"))
        return dest

    def move(self, bucket: str, src: str, dst: str) -> None:
        body = json.dumps({"bucketId": bucket, "sourceKey": src, "destinationKey": dst}).encode()
        self._request("POST", "/storage/v1/object/move", body, {"Content-Type": "application/json"})


def find_or_create(sb: Supabase, table: str, match: dict, extra: dict) -> dict:
    rows = sb.select(table, {"select": "*", **{k: ("is.null" if v is None else f"eq.{v}") for k, v in match.items()}})
    if rows:
        return rows[0]
    return sb.insert(table, {**match, **extra})


def ensure_course(sb: Supabase, centre_id: str, f: LessonFile) -> str:
    level_name, stage, sort = (f.level_name, *_level_meta(f.level_code))
    # Prefer platform-wide catalogue rows; otherwise create the centre's own.
    def catalogue(table: str, code: str, extra: dict) -> dict:
        shared = sb.select(table, {"select": "*", "code": f"eq.{code}", "centre_id": "is.null"})
        if shared:
            return shared[0]
        return find_or_create(sb, table, {"centre_id": centre_id, "code": code}, extra)

    subject = catalogue("subjects", f.subject_code, {"name": f.subject_name})
    level = catalogue("levels", f.level_code, {"name": level_name, "stage": stage, "sort": sort})
    for centre_filter in ("is.null", f"eq.{centre_id}"):
        rows = sb.select("courses", {"select": "id", "subject_id": f"eq.{subject['id']}", "level_id": f"eq.{level['id']}", "centre_id": centre_filter})
        if rows:
            return rows[0]["id"]
    return sb.insert("courses", {"centre_id": centre_id, "subject_id": subject["id"], "level_id": level["id"], "name": f.course_name})["id"]


def _level_meta(code: str) -> tuple[str, int]:
    from naming import LEVEL_NAMES

    _, stage, sort = LEVEL_NAMES[code]
    return stage, sort


def publish_lesson(sb: Supabase, centre_id: str, files: dict[str, tuple[LessonFile, Path]], workdir: Path, publish: bool) -> str:
    first = next(iter(files.values()))[0]
    course_id = ensure_course(sb, centre_id, first)
    title = first.title
    lessons = sb.select("course_lessons", {"select": "id", "centre_id": f"eq.{centre_id}", "course_id": f"eq.{course_id}", "position": f"eq.{first.week}"})
    values = {"title": title, "status": "published" if publish else "draft"}
    if lessons:
        lesson_id = lessons[0]["id"]
        sb.update("course_lessons", {"id": lesson_id}, values)
    else:
        lesson_id = sb.insert("course_lessons", {"centre_id": centre_id, "course_id": course_id, "position": first.week, **values})["id"]

    report = []
    for part_name, (info, source) in sorted(files.items(), key=lambda kv: PART_SETTINGS[kv[0]]["position"]):
        built = build_part(info, source, workdir / part_name)
        prefix = f"centres/{centre_id}/lessons/{lesson_id}/{part_name}"
        pdf_path = f"{prefix}.pdf"
        sb.upload("lesson-files", pdf_path, built.pdf, "application/pdf")
        pages = []
        for n, page in enumerate(built.pages, start=1):
            path = f"{prefix}/{n:03d}.webp"
            sb.upload("lesson-files", path, page.image, "image/webp")
            entry = {"path": path, "role": page.role}
            if page.title:
                entry["title"] = page.title
            pages.append(entry)

        settings = PART_SETTINGS[part_name]
        row = {
            "kind": "document",
            "title": settings["title"],
            "file_path": pdf_path,
            "pages": pages,
            "downloadable": settings["downloadable"],
        }
        existing = sb.select("lesson_parts", {"select": "id", "centre_id": f"eq.{centre_id}", "lesson_id": f"eq.{lesson_id}", "position": f"eq.{settings['position']}"})
        if existing:
            sb.update("lesson_parts", {"id": existing[0]["id"]}, row)
        else:
            sb.insert("lesson_parts", {"centre_id": centre_id, "lesson_id": lesson_id, "position": settings["position"], **row})
        solutions = sum(1 for p in built.pages if p.role == "solution")
        hidden = f", hid {len(built.hidden)} ({'; '.join(built.hidden)})" if built.hidden else ""
        report.append(f"  {settings['title']}: {len(built.pages)} pages, {solutions} worked solutions{hidden}")

    print(f"✓ {first.course_name} · Week {first.week}: {title} ({'published' if publish else 'draft'})")
    print("\n".join(report))
    return lesson_id


# --------------------------------------------------------------------------------------
# Commands
# --------------------------------------------------------------------------------------
def cmd_build(args: argparse.Namespace) -> int:
    out = Path(args.out)
    lessons, problems = group_files([Path(p) for p in args.files])
    for p in problems:
        print("! " + p)
    for key, files in lessons.items():
        for part_name, (info, source) in files.items():
            built = build_part(info, source, out / "-".join(map(str, key)) / part_name)
            manifest = [{"image": p.image.name, "role": p.role, "title": p.title} for p in built.pages]
            (out / "-".join(map(str, key)) / f"{part_name}.json").write_text(json.dumps({"hidden": built.hidden, "pages": manifest}, ensure_ascii=False, indent=1))
            print(f"✓ {info.course_name} · Week {info.week} · {part_name}: {len(built.pages)} pages "
                  f"({sum(p.role == 'solution' for p in built.pages)} solutions, hidden: {built.hidden or 'none'})")
    return 1 if problems and not lessons else 0


def cmd_inbox(args: argparse.Namespace) -> int:
    url, key = os.environ.get("SUPABASE_URL"), os.environ.get("SUPABASE_SERVICE_ROLE_KEY")
    if not url or not key:
        print("SUPABASE_URL and SUPABASE_SERVICE_ROLE_KEY must be set.", file=sys.stderr)
        return 2
    sb = Supabase(url, key)
    centres = sb.select("centres", {"select": "id,name", "slug": f"eq.{args.centre}"})
    if not centres:
        print(f"No centre with slug '{args.centre}'.", file=sys.stderr)
        return 2
    centre_id = centres[0]["id"]

    names = [n for n in sb.list_inbox("lesson-inbox") if "/" not in n]
    if not names:
        print("The lesson-inbox bucket is empty. Upload files there first.")
        return 0

    stamp = time.strftime("%Y%m%d-%H%M%S")
    failures = 0
    with tempfile.TemporaryDirectory() as tmp:
        tmpdir = Path(tmp)
        local = [sb.download("lesson-inbox", n, tmpdir / n) for n in names]
        lessons, problems = group_files(local)
        for p in problems:
            print("! " + p)
        for key, files in lessons.items():
            try:
                publish_lesson(sb, centre_id, files, tmpdir / "work" / "-".join(map(str, key)), publish=not args.draft)
                for _, source in files.values():
                    sb.move("lesson-inbox", source.name, f"done/{stamp}/{source.name}")
            except Exception as err:  # keep going with the other lessons
                failures += 1
                print(f"✗ {'-'.join(map(str, key))}: {err}")
                for _, source in files.values():
                    sb.move("lesson-inbox", source.name, f"failed/{stamp}/{source.name}")
            finally:
                shutil.rmtree(tmpdir / "work", ignore_errors=True)
    return 1 if failures else 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="command", required=True)
    b = sub.add_parser("build", help="convert and render locally, no upload")
    b.add_argument("--out", required=True)
    b.add_argument("files", nargs="+")
    i = sub.add_parser("inbox", help="import everything in the lesson-inbox bucket")
    i.add_argument("--centre", required=True, help="centre slug, e.g. aspire-hub")
    i.add_argument("--draft", action="store_true", help="import as drafts (hidden from students)")
    args = parser.parse_args()
    return cmd_build(args) if args.command == "build" else cmd_inbox(args)


if __name__ == "__main__":
    sys.exit(main())
